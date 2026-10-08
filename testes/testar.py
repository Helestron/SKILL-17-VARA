#!/usr/bin/env python3
"""Testes de ponta a ponta dos scripts da skill (sem rede e sem o Helestron instalado).

Uso: python3 testes/testar.py   (requer reportlab e pdfplumber ou pypdf para a pasta sintética)
"""
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
SK = RAIZ / "skill-17a-lote-minutas-helestron"
S = SK / "scripts"
PY = sys.executable
falhas = []


def rodar(*args, esperado=0):
    r = subprocess.run([PY, "-I", *map(str, args)], capture_output=True, text=True, encoding="utf-8")
    if esperado is not None and r.returncode != esperado:
        falhas.append(f"{' '.join(map(str, args[:3]))}: código {r.returncode} (esperado {esperado})\n{r.stdout[-800:]}{r.stderr[-800:]}")
    return r


def checar(cond, msg):
    if not cond:
        falhas.append(msg)


def subst(texto: str, a: str, b: str) -> str:
    assert a in texto, f"trecho ausente no texto-base do teste: {a!r}"
    return texto.replace(a, b)


def minuta(tmp: Path, nome: str, texto: str) -> Path:
    p = tmp / nome
    p.write_text(texto, encoding="utf-8")
    return p


def main():
    tmp = Path(tempfile.mkdtemp(prefix="vara17_"))
    try:
        # 1. pasta sintética do Helestron, inventário, texto, transcrições e sigilo
        n = json.loads(subprocess.run([PY, "-I", RAIZ / "testes" / "criar_fixture.py", tmp / "acervo", S],
                                      capture_output=True, text=True, check=True).stdout)
        r = rodar(S / "autos.py", "inventario", tmp / "acervo", "--lista", n["n1"][:15], n["n2"], n["n3"], "0000001-00.2024.8.02.0001")
        checar("TRABALHO=" in r.stdout and "inválidos" in r.stdout and "(processo sigiloso)" in r.stdout, "inventário")
        checar(n["n3"] not in r.stdout, "número do sigiloso exposto no inventário")
        t = r.stdout.split("TRABALHO=")[1].splitlines()[0].strip()
        r = rodar(S / "autos.py", "preparar", "-t", t)
        checar("1 transcrição" in r.stdout and "Mandado de Segurança" in r.stdout, "preparar")
        r = rodar(S / "autos.py", "ler", "1", "--fls", "4-4", "-t", t)
        checar("[fl. 4]" in r.stdout and "CONTESTACAO" in r.stdout, "ler por folhas")
        r = rodar(S / "autos.py", "transcricoes", "1", "--ler", "1", "-t", t)
        checar("[12:30] Testemunha" in r.stdout and "2024-05-10" in r.stdout, "transcrição .srt")
        rodar(S / "autos.py", "ler", "3", "--fls", "1-1", "-t", t, esperado=1)  # sigiloso sem autorização
        rodar(S / "autos.py", "autorizar", "3", "--ordem", "autorizo o sigiloso", "-t", t)
        rodar(S / "autos.py", "preparar", "3", "-t", t)
        r = rodar(S / "autos.py", "ler", "3", "--fls", "1-1", "-t", t)
        checar("SEGREDO" in r.stdout, "sigiloso autorizado")
        r = rodar(S / "autos.py", "inventario", tmp / "acervo", "--lista", n["n1"], n["n2"], n["n3"])
        checar("lote01" in r.stdout, "a mesma lista deve retomar o lote aberto")

        # 1-b. transcrições em .vtt e .docx (formatos alternativos do aplicativo)
        tr = tmp / "Transcrições"
        (tr / f"{n['n2']}_audiencia_15-03-2024.vtt").write_text(
            "WEBVTT\n\n00:01:05.000 --> 00:01:09.000\nJuiz: Ouvida a testemunha.\n", encoding="utf-8")
        with zipfile.ZipFile(tr / f"Termo {n['n2']}.docx", "w") as z:
            z.writestr("word/document.xml", '<w:document xmlns:w="x"><w:body><w:p><w:r><w:t>Depoimento: o fiscal reteve a carga.</w:t></w:r></w:p></w:body></w:document>')
        r = rodar(S / "ponte_helestron.py", "transcricoes", "--transcricoes", tr, "--numero", n["n2"])
        idx = json.loads(r.stdout)["indice"][n["n2"]]
        checar(len(idx) == 2 and idx[0]["data"] == "2024-03-15", "índice de transcrições .vtt/.docx")
        r = rodar(S / "ponte_helestron.py", "ler-transcricao", next(x["arquivo"] for x in idx if x["tipo"] == "vtt"))
        checar("[01:05] Juiz: Ouvida a testemunha." in r.stdout, "leitura de .vtt")
        r = rodar(S / "ponte_helestron.py", "ler-transcricao", next(x["arquivo"] for x in idx if x["tipo"] == "docx"))
        checar("o fiscal reteve a carga" in r.stdout, "leitura de .docx")

        # 2. portão
        exemplo = RAIZ / "testes" / "minuta_exemplo_sentenca.txt"
        r = rodar(S / "verificar_minuta.py", exemplo)
        checar(json.loads(r.stdout)["resultado"] == "OK", "exemplo deveria passar no portão")
        base = exemplo.read_text(encoding="utf-8")
        ruins = {
            "numerada": subst(base, "Réplica às fls. 70/75.", "5. Réplica às fls. 70/75."),
            "primeira_pessoa": subst(base, "Com efeito, o autor", "Entendo que o autor"),
            "vedado": subst(base, "Com efeito,", "Compulsando os autos,"),
            "sem_pri": subst(base, "P. R. I.", "Publique-se."),
            "titulo": subst(base, "É o Relatório.\n", "É o Relatório.\nDa preliminar\n"),
            "metodo": subst(base, "Com efeito,", "Feita a leitura integral dos autos,"),
            "desbalanceada": subst(base, "**coisa julgada**", "**coisa julgada"),
        }
        for nome, txt in ruins.items():
            r = rodar(S / "verificar_minuta.py", minuta(tmp, f"{nome}.txt", txt), esperado=1)
            checar(json.loads(r.stdout)["resultado"] == "BLOQUEADO", f"portão deveria bloquear: {nome}")

        # 3. geração do Word
        saida = tmp / "minutas"
        r = rodar(S / "gerar_minuta.py", exemplo, "--saida", saida)
        res = json.loads(r.stdout)
        checar(res["resultado"] == "OK", "geração")
        limpa, anot = Path(res["limpa"]), Path(res["anotada"])
        xl = zipfile.ZipFile(limpa).read("word/document.xml").decode()
        xa = zipfile.ZipFile(anot).read("word/document.xml").decode()
        checar("FF0000" not in xl and "Conferir" not in xl and "<w:numPr>" not in xl, "limpa sem vermelho/numeração")
        checar("FF0000" in xa and "Conferir" in xa and "Resolução CNJ" in xa, "anotada com vermelho e ressalva")
        checar("Courier New" in xl and 'w:ind w:left="2268"' in xl and 'w:firstLine="1417"' in xl, "formatação dos modelos")
        rodar(S / "gerar_minuta.py", minuta(tmp, "bloq.txt", ruins["numerada"]), "--saida", saida, esperado=1)
        checar(not (saida / "bloq.docx").exists(), "minuta bloqueada não pode gerar .docx")

        # 4. ledger
        r = rodar(S / "ledger.py", "conferir", minuta(tmp, "cit.txt", subst(base, "(art. 373, I, do CPC)", "(art. 373, I, do CPC; REsp n.º 1.340.553/RS)")), "-t", t, esperado=1)
        checar("REsp n.º 1.340.553" in r.stdout, "precedente sem ledger deve bloquear")
        e = [{"tipo": "precedente", "chave": "REsp n.º 1.340.553", "tribunal": "STJ", "orgao": "Primeira Seção", "relator": "Min. X",
              "julgamento": "2018-09-12", "fonte": "https://scon.stj.jus.br/x", "status": "VERIFIED", "aderencia": "teste",
              "trecho_literal": "RECURSO ESPECIAL REPETITIVO. PRESCRIÇÃO INTERCORRENTE. ART. 40 DA LEF."},
             {"tipo": "precedente", "chave": "IRDR n.º 3", "tribunal": "TJAL", "orgao": "Seção Especializada Cível",
              "relator": "Des. Y", "julgamento": "2025-07-08", "fonte": "https://www2.tjal.jus.br/cjsg/x", "status": "VERIFIED",
              "aderencia": "teste", "trecho_literal": "1.1. O mero cumprimento do interstício temporal no posto não implica preterição."}]
        (tmp / "e.json").write_text(json.dumps(e), encoding="utf-8")
        rodar(S / "ledger.py", "add", tmp / "e.json", "-t", t)
        rodar(S / "ledger.py", "conferir", tmp / "cit.txt", "-t", t, esperado=0)
        e[0]["fonte"] = "https://blog.exemplo.com/x"
        (tmp / "e2.json").write_text(json.dumps(e), encoding="utf-8")
        rodar(S / "ledger.py", "add", tmp / "e2.json", "-t", t, esperado=1)
        rodar(S / "ledger.py", "calc", tmp / "c.json", "add", "--descricao", "x", "--expr", "__import__('os')", esperado=1)

        # 5. modelos
        for m in sorted((SK / "modelos").glob("*.txt")):
            txt = m.read_text(encoding="utf-8")
            checar(txt.splitlines()[1].startswith("@ato "), f"modelo sem @ato: {m.name}")
        r = rodar(S / "importar_modelo.py", "buscar", "promoção militar inativo")
        checar("inativo" in r.stdout, "busca de modelos")
        if shutil.which("soffice"):
            r = rodar(S / "importar_modelo.py", next((SK / "modelos" / "originais").glob("*0750345*")))
            checar("@ato sentença" in r.stdout and "!! Desse modo, determino a exclusão" in r.stdout, "importação de RTF")

        # 6. regressões da revisão adversarial
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "afl.txt", subst(base, "Réplica às fls. 70/75.", "Réplica às fls. 70/75, com documento juntado à fl. 76.")))
        checar(json.loads(r.stdout)["resultado"] == "OK", "'à fl.' é a forma correta e não pode bloquear")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "semproc.txt", subst(base, "@processo 0714346-41.2024.8.02.0001\n", "")), esperado=1)
        checar("@processo" in r.stdout, "@processo obrigatório")
        desp = ("@ato despacho\n@processo 0714346-41.2024.8.02.0001\nDefiro a produção de prova documental requerida às fls. 80/81.\n"
                "Intime-se o Estado de Alagoas para juntar a ficha funcional do autor no prazo de 15 (quinze) dias.\n"
                "Decorrido o prazo, voltem conclusos para sentença.\nCumpra-se.\n")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "desp.txt", desp))
        checar(json.loads(r.stdout)["resultado"] == "OK", "despacho com 1ª pessoa deve passar")
        r = rodar(S / "gerar_minuta.py", tmp / "desp.txt", "--saida", saida)
        xd = zipfile.ZipFile(json.loads(r.stdout)["limpa"]).read("word/document.xml").decode()
        checar(xd.count('<w:u w:val="single"/>') >= 2, "comandos do despacho sublinhados")
        homol = subst(base, "Diante do exposto, julgo improcedente a demanda.", "Diante do exposto, homologo o acordo de fls. 90/92.")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "homol.txt", homol))
        checar(json.loads(r.stdout)["resultado"] == "OK", "sentença homologatória deve passar")
        r = rodar(S / "ledger.py", "calc", tmp / "c2.json", "add", "--descricao", "razão", "--expr", "round(4.908,00 / 1.412,00, 2)")
        checar('"resultado": 3.48' in r.stdout, "cálculo com decimal brasileiro e round(x, 2)")
        r = rodar(S / "autos.py", "inventario", tmp / "acervo", "--lista", n["n1"], "-t", tmp / "outro_trab")
        checar("outro_trab" in r.stdout and "inválidos" not in r.stdout, "-t depois de --lista")
        r = rodar(S / "autos.py", "marcar", "1", "minutado", "--justificativa", "gerúndio mantido: citação", "-t", tmp / "outro_trab")
        r = rodar(S / "autos.py", "marcar", "1", "minutado", "--justificativa", "dois-pontos mantido: bloco", "-t", tmp / "outro_trab")
        est = json.loads((tmp / "outro_trab" / "estado.json").read_text(encoding="utf-8"))
        checar("gerúndio" in est["processos"][n["n1"]]["justificativa"] and "dois-pontos" in est["processos"][n["n1"]]["justificativa"],
               "justificativas acumuladas")
        verm = subst(base, "Diante do exposto, julgo improcedente a demanda.",
                            "Diante do exposto, julgo {{Conferir: pedido subsidiário}} improcedente a demanda.")
        r = rodar(S / "gerar_minuta.py", minuta(tmp, "verm.txt", subst(verm, "@processo 0714346-41.2024.8.02.0001", "@processo 0750345-89.2023.8.02.0001")), "--saida", saida)
        res = json.loads(r.stdout)
        checar(res["resultado"] == "OK", "apontamento no início do dispositivo não pode reprovar a limpa")
        xv = zipfile.ZipFile(res["limpa"]).read("word/document.xml").decode()
        texto_v = "".join(__import__("re").findall(r"<w:t[^>]*>([^<]*)</w:t>", xv))
        checar("julgo improcedente" in texto_v and "  " not in texto_v, "limpa sem espaço duplo")
        serie = subst(base, "> 1.1. O mero", "> Apelação Cível n.º 0702075-63.2025.8.02.0001, Rel. Des. A, julgado em 12/08/2026.\n"
                             "> Apelação Cível n.º 0726628-14.2024.8.02.0001, Rel. Des. A, julgado em 12/08/2026.\n"
                             "> Apelação Cível n.º 0739309-84.2022.8.02.0001, Rel. Des. B, julgado em 05/08/2026.\n> 1.1. O mero")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "serie.txt", serie))
        checar("precedentes em série" in r.stdout, "precedentes em série")
        # Helestron simulado: sigilosos fora do acervo, pelo `caminhos --json`
        fora = tmp / "SigilososFora" / "Lote X"
        fora.mkdir(parents=True)
        shutil.copy(next((tmp / "acervo").glob(f"{n['n2']}.pdf")), fora / "0700099-55.2024.8.02.0001.pdf")
        falso = tmp / "helestron_falso"
        falso.write_text("#!/bin/sh\necho '{\"versao\": \"1.0.9\", \"acervo\": \"" + str(tmp / "acervo") + "\", \"sigilosos\": \""
                         + str(tmp / "SigilososFora") + "\"}'\n", encoding="utf-8")
        falso.chmod(0o755)
        env = dict(__import__("os").environ, HELESTRON_PYTHON=str(falso))
        r = subprocess.run([PY, "-I", str(S / "autos.py"), "inventario", "-t", str(tmp / "trab_sig")], capture_output=True,
                           text=True, encoding="utf-8", env=env)
        checar("0700099-55.2024" not in r.stdout and r.stdout.count("(processo sigiloso)") >= 2,
               "sigiloso fora do acervo deve entrar no inventário, sem expor o número")
        # PDF de outra origem com carimbo "fls. 100": a folha é a carimbada
        from reportlab.pdfgen import canvas
        outra = tmp / "outra"; outra.mkdir()
        c = canvas.Canvas(str(outra / "0700100-40.2024.8.02.0001.pdf"))
        for k in range(3):
            c.drawString(500, 810, f"fls. {100 + k}"); c.drawString(60, 780, f"Peça de outra origem, página {k + 1}"); c.showPage()
        c.save()
        r = rodar(S / "autos.py", "inventario", outra, "-t", tmp / "trab_outra")
        rodar(S / "autos.py", "preparar", "-t", tmp / "trab_outra")
        r = rodar(S / "autos.py", "ler", "1", "--fls", "101-101", "-t", tmp / "trab_outra")
        checar("[fl. 101]" in r.stdout and "página 2" in r.stdout, "folha carimbada prevalece sobre a posição no PDF")

        # 7. cumprimento de sentença (versão 2.1)
        cs = RAIZ / "testes" / "minuta_exemplo_despacho_cs.txt"
        r = rodar(S / "verificar_minuta.py", cs)
        res = json.loads(r.stdout)
        checar(res["resultado"] == "OK", f"despacho de cumprimento de exemplo deveria passar: {res['pendencias_bloqueantes']}")
        checar(not any("Contadoria" in a for a in res["apontamentos"]), "'não é cabível o encaminhamento à Contadoria' não é remessa")
        r = rodar(S / "gerar_minuta.py", cs, "--saida", saida)
        xc = zipfile.ZipFile(json.loads(r.stdout)["limpa"]).read("word/document.xml").decode()
        import re as _re
        itens = [p for p in _re.findall(r"<w:p>.*?</w:p>", xc, _re.S) if 'w:ind w:left="2268"/>' in p or 'w:ind w:left="2551"/>' in p]
        checar(len(itens) == 5, f"cinco itens recuados (achados {len(itens)})")
        checar(all('<w:u w:val="single"/>' in p for p in itens), "itens do dispositivo sublinhados")
        checar(any('w:ind w:left="2551"/>' in p for p in itens), "subitens com recuo de 4,5 cm")
        base_cs = cs.read_text(encoding="utf-8")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "cs_item.txt", subst(base_cs, "Decorrido o prazo fixado para a parte exequente",
                                                                              "Decorrido o prazo fixado no item 4 para a parte exequente")))
        checar("remissão a 'item N'" in r.stdout, "remissão a item N apontada")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "cs_cont.txt", subst(base_cs, "Cumpra-se observada a sequência acima.",
                                                                              "Remetam-se os autos à Contadoria.\nCumpra-se.")))
        checar("Contadoria só se não houver meio eletrônico" in r.stdout, "remessa à Contadoria apontada")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "cs_fichas.txt", subst(base_cs, "Cumpra-se observada a sequência acima.",
                                                                                "Intime-se o Estado de Alagoas para juntar as fichas financeiras do exequente.\nCumpra-se.")))
        checar("ônus do exequente" in r.stdout, "fichas financeiras impostas ao executado apontadas")
        sent_cs = subst(subst(base, "Trata-se de Ação Ordinária proposta", "Trata-se de Cumprimento de Sentença proposto"),
                        "Condeno a parte autora nas custas e em honorários", "Condeno a parte executada nas custas e em honorários")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "cs_custas.txt", sent_cs))
        checar("não há custas" in r.stdout, "custas no cumprimento de sentença apontadas")
        for m in ("despacho_CS_adequar_requerimento_art534_Res21-2023_espolio_sem_inventariante.txt",
                  "despacho_CS_habilitacao_sem_inventario_intima_herdeiras_indicar_endereco_de_outra_herdeira.txt",
                  "despacho_CS_fichas_financeiras_onus_da_exequente.txt",
                  "despacho_CS_cessao_de_credito_de_precatorio_Presidencia_TJ_reitera_Secretaria_nao_volta_concluso.txt"):
            txt = (SK / "modelos" / m).read_text(encoding="utf-8")
            checar("@ato despacho" in txt and "// rodapé: Maceió" in txt and "\n> " not in txt, f"modelo CS importado: {m}")
        txt = (SK / "modelos" / "despacho_CS_adequar_requerimento_art534_Res21-2023_espolio_sem_inventariante.txt").read_text(encoding="utf-8")
        checar(txt.count("\n+ ") == 6 and txt.count("\n++ ") == 4 and "\n5. " not in txt, "enumeração recuada e sem numeração digitada")
        r = rodar(S / "importar_modelo.py", "buscar", "cumprimento habilitação herdeira")
        checar("despacho_CS_habilitacao" in r.stdout, "busca de modelo de cumprimento")
        # autos relacionados: incidente -01 com o principal na pasta
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(tmp / "acervo" / f"{n['n1']}-01.pdf"))
        for k, linha in enumerate(["CUMPRIMENTO DE SENTENCA. Planilha de calculo com IPCA-E.", "Certidao de obito. Espolio. Herdeiros.",
                                   "Contrato de honorarios para destaque. Dados bancarios."]):
            c.drawString(500, 810, f"fls. {k + 1}"); c.drawString(40, 780, linha); c.showPage()
        c.save()
        r = rodar(S / "autos.py", "inventario", tmp / "acervo", "--lista", n["n1"][:15] + "/01", "-t", tmp / "trab_cs")
        checar(f"apoio: {n['n1']}" in r.stdout and "cumprimento de sentença" in r.stdout, "incidente aponta os autos de origem")
        checar("prioridade" not in r.stdout, "incidente não herda a capa do principal")
        r = rodar(S / "autos.py", "preparar", "-t", tmp / "trab_cs")
        checar("(apoio)" in r.stdout, "preparar extrai os autos de apoio")
        r = rodar(S / "autos.py", "requisitos", "1", "-t", tmp / "trab_cs")
        checar("planilha ou memória de cálculo: fl. 1" in r.stdout and "óbito, espólio" in r.stdout and
               "contrato de honorários (destaque): fl. 3" in r.stdout, "checklist de requisitos do cumprimento")
        r = rodar(S / "autos.py", "inventario", tmp / "acervo", "--lista", "0750345-89.2023/02", "-t", tmp / "trab_cs2")
        checar("não encontrados" in r.stdout, "incidente ausente relatado")

        # 8. regressões da revisão das adições de cumprimento (versão 2.1)
        import re as _re
        desp2 = ("@ato despacho\n@processo 0714346-41.2024.8.02.0001-01\n"
                 "Havendo nos autos prova do óbito, a habilitação depende da participação de todos os herdeiros.\n"
                 "Havendo impugnação, intime-se a parte exequente para manifestação em 15 (quinze) dias.\nCumpra-se.\n")
        r = rodar(S / "gerar_minuta.py", minuta(tmp, "desp2.txt", desp2), "--saida", saida)
        x2 = zipfile.ZipFile(json.loads(r.stdout)["limpa"]).read("word/document.xml").decode()
        ps2 = _re.findall(r"<w:p>.*?</w:p>", x2, _re.S)
        checar('<w:u ' not in ps2[0] and '<w:u ' in ps2[1], "abertura condicional só é comando com verbo de comando")
        mod = SK / "modelos" / "despacho_CS_adequar_requerimento_art534_Res21-2023_espolio_sem_inventariante.txt"
        r = rodar(S / "gerar_minuta.py", mod, "--saida", saida, "--nome", "modelo_adequar", "--rascunho", esperado=None)
        xm = zipfile.ZipFile(saida / "modelo_adequar_anotada.docx").read("word/document.xml").decode()
        itens_m = [p for p in _re.findall(r"<w:p>.*?</w:p>", xm, _re.S) if 'w:ind w:left="2268"/>' in p or 'w:ind w:left="2551"/>' in p]
        checar(len(itens_m) == 10 and all('<w:u ' in p for p in itens_m), "itens do modelo de adequação sublinhados como no original")
        for nome, txt, esperado in (
                ("registre_oportuno", subst(base, "Com efeito, o autor", "Registre-se, por oportuno, que o autor"), "OK"),
                ("pub_int", subst(base, "P. R. I.", "Publique-se e intimem-se."), "BLOQUEADO")):
            r = rodar(S / "verificar_minuta.py", minuta(tmp, nome + ".txt", txt), esperado=None)
            checar(json.loads(r.stdout)["resultado"] == esperado, f"fórmula de publicação: {nome}")
        conh = subst(base, "Réplica às fls. 70/75.", "Réplica às fls. 70/75, com valores a apurar em cumprimento de sentença.")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "conh.txt", conh))
        checar("cumprimento de sentença: não há custas" not in r.stdout, "sentença de conhecimento não é tratada como cumprimento")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "edital.txt", subst(base_cs, "tornem os autos imediatamente conclusos.",
                                                                             "tornem os autos imediatamente conclusos, observado o item 10 do edital.")))
        checar("remissão a 'item N'" not in r.stdout, "item de edital não é remissão interna")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "remessa.txt", subst(base_cs, "Cumpra-se observada a sequência acima.",
                                                                              "Determino a remessa dos autos à Contadoria Judicial.\nCumpra-se.")))
        checar("Contadoria só se não houver meio eletrônico" in r.stdout, "'determino a remessa … à Contadoria' apontada")
        r = rodar(S / "verificar_minuta.py", minuta(tmp, "oficio.txt", subst(base_cs, "Cumpra-se observada a sequência acima.",
                                                                             "Oficie-se à Secretaria de Estado da Fazenda para enviar as fichas financeiras do exequente.\nCumpra-se.")))
        checar("ônus do exequente" in r.stdout, "ofício ao órgão pagador pedindo fichas financeiras apontado")
        # importador: fecho exato, transcrição recuada em Times, numeração dentro do negrito, cabeçalho após o título
        import docx as _docx
        from docx.shared import Cm
        d = _docx.Document()
        for linha in ("D E S P A C H O", "AUTOR: FULANO DE TAL"):
            d.add_paragraph(linha)
        d.add_paragraph("Trata-se de Cumprimento de Sentença proposto por Fulano em face do Estado de Alagoas.")
        d.add_paragraph("Cumpra-se a decisão de fls. 80.")
        pr = d.add_paragraph("O art. 534 do CPC dispõe que o exequente apresentará demonstrativo discriminado e atualizado do crédito.")
        pr.paragraph_format.left_indent = Cm(4)
        p7 = d.add_paragraph(); rr = p7.add_run("7. Não havendo impugnação"); rr.bold = True; p7.add_run(" ou havendo concordância, conclusos.")
        d.add_paragraph("Cumpra-se.")
        d.add_paragraph("Maceió, datado eletronicamente.")
        d.save(str(tmp / "imp.docx"))
        r = rodar(S / "importar_modelo.py", tmp / "imp.docx")
        checar("// cabeçalho: AUTOR: FULANO DE TAL" in r.stdout, "linha de cabeçalho após o título fica como comentário")
        checar("\nCumpra-se a decisão de fls. 80.\n" in r.stdout, "'Cumpra-se a decisão…' não encerra o corpo")
        checar("\n> O art. 534 do CPC" in r.stdout, "transcrição recuada em Times continua transcrição")
        checar("\n**Não havendo impugnação** ou havendo" in r.stdout, "numeração dentro do negrito sem desbalancear")
        # pasta por processo: o principal não herda a capa do incidente
        pp = tmp / "porprocesso" / n["n1"]
        (pp / "_controle").mkdir(parents=True)
        shutil.copy(next((tmp / "acervo").glob(f"{n['n1']}.pdf")), pp / "autos.pdf")
        (pp / "_controle" / f"{n['n1']}-01_capa.json").write_text(json.dumps({"capa": {"classe": "Cumprimento de Sentença"}}), encoding="utf-8")
        (pp / "_controle" / f"{n['n1']}_capa.json").write_text(json.dumps({"capa": {"classe": "Procedimento Comum Cível"}}), encoding="utf-8")
        r = rodar(S / "autos.py", "inventario", tmp / "porprocesso", "-t", tmp / "trab_pp")
        checar("Procedimento Comum Cível" in r.stdout, "pasta por processo: o principal lê a própria capa")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if falhas:
        print("FALHAS:\n- " + "\n- ".join(falhas))
        return 1
    print("todos os testes passaram")
    return 0


if __name__ == "__main__":
    sys.exit(main())
